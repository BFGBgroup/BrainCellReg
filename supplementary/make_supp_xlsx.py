# -*- coding: utf-8 -*-
"""Supplementary Data 1-4 as Excel files (Nat Commun: one file per Supplementary Data, a Description sheet first).
Only renames columns and disease codes; no number is changed.
  Data 1  Fisher enrichment, all disease x context tests       <results>/Supplementary_Data_1_Fisher.tsv
  Data 2  SMR, HEIDI and colocalization, 2,578 combinations     <results>/Supplementary_Data_2_SMR_coloc.tsv
  Data 3  scGWAS significant modules (118)                      <scGWAS / external SMR tables> (analysis/s82_supp_data.py)
  Data 4  external SMR replication, 253 pairs + per dataset     same
python make_supp_xlsx.py <out_dir> <version>"""
import sys, io, os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
out, ver = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
MS = r'<path to the Fisher and SMR result tables>'
SV = r'<path to the scGWAS and external SMR tables>'
DCODE = {'MDD': 'DEP', 'CUD': 'CD', 'LacunarStroke': 'LS', 'FebrileSeizures': 'FS', 'Narcolepsy': 'NAR'}
DIS_FULL = ('MG, myasthenia gravis; MS, multiple sclerosis; AD, Alzheimer\'s disease; SCZ, schizophrenia; DEP, depression; '
            'LS, lacunar stroke; PD, Parkinson\'s disease; FS, febrile seizures; NAR, narcolepsy; CD, cannabis dependence.')
CTX = 'Context: brain region and deconvolved cell type, written REGION.CellType (for example, FCTX.ExN).'


def write(path, title, notes, cols, sheets):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Description'
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=12)
    r = 3
    for n in notes:
        ws.cell(r, 1, n).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
    r += 1
    for sheet, cdesc in cols:
        ws.cell(r, 1, f'Columns of sheet "{sheet}"').font = Font(bold=True)
        r += 1
        for c, d in cdesc:
            ws.cell(r, 1, c).font = Font(name='Consolas')
            ws.cell(r, 2, d).alignment = Alignment(wrap_text=True, vertical='top')
            r += 1
        r += 1
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 110
    for name, df in sheets:
        s = wb.create_sheet(name)
        s.append(list(df.columns))
        for c in s[1]:
            c.font = Font(bold=True)
        for row in df.itertuples(index=False):
            s.append([None if (isinstance(v, float) and pd.isna(v)) else (v.item() if hasattr(v, 'item') else v) for v in row])
        s.freeze_panes = 'A2'
        for i, c in enumerate(df.columns, 1):
            s.column_dimensions[get_column_letter(i)].width = min(45, max(10, len(str(c)) + 2))
        documented = [c for c, _ in dict(cols)[name]]
        assert documented == list(df.columns), (name, set(df.columns) ^ set(documented))
    wb.save(path)
    print(path, [(n, len(d)) for n, d in sheets])


# ---- Data 1
d1 = pd.read_csv(f'{MS}/Supplementary_Data_1_Fisher.tsv', sep='\t')
c1 = [('disease_group', 'Neuropsychiatric (ICD-11 chapters 6-8; 94 disorders) or other ICD-11 chapters (417 diseases).'),
      ('disease', 'Disease name.'),
      ('context', CTX), ('region', 'Brain region.'), ('cell_type', 'Deconvolved cell type.'),
      ('eSNP_associated', 'Independent eSNPs of the context that were disease-associated variants.'),
      ('eSNP_not_associated', 'Independent eSNPs that were not disease-associated.'),
      ('non_eSNP_associated', 'Non-eSNPs (after LD clumping) that were disease-associated.'),
      ('non_eSNP_not_associated', 'Non-eSNPs that were not disease-associated.'),
      ('OR', 'Odds ratio from two-sided Fisher\'s exact test.'), ('OR_lower95', 'Lower 95% confidence limit of the OR.'),
      ('OR_upper95', 'Upper 95% confidence limit of the OR.'), ('P', 'Fisher\'s exact test P.'),
      ('q', 'Benjamini-Hochberg q across all tests of the disease set (neuropsychiatric or other).'),
      ('significant', 'TRUE when q <= 0.05, OR_lower95 > 1, all four cells > 0, and the test was not removed as unstable (OR > 1,000 with fewer than three variants in the eSNP-unassociated or non-eSNP-associated cell).'),
      ('OR_noMHC', 'OR after removing all variants in chr6:25-35 Mb (hg38) from the eSNP, non-eSNP and disease sets. The four noMHC columns are empty where this test was not run (the 511 rows of TCTX.OPC).'),
      ('P_noMHC', 'P of the test without the MHC region.'), ('q_noMHC', 'q of the test without the MHC region.'),
      ('significant_noMHC', 'Significance of the test without the MHC region, same rule as significant.')]
write(f'{out}/Supplementary_Data_1_v{ver}.xlsx', 'Supplementary Data 1. Enrichment of disease-associated variants among eSNPs in each context',
      ['One row per disease and context (Fisher\'s exact test). Disease-associated variants were the variants reported for each disease in the GWAS Catalog, extended to variants in LD (r2 >= 0.8 within 1 Mb, 1000 Genomes European samples).',
       'Contexts without independent eSNPs were not tested. ' + CTX],
      [('Fisher', c1)], [('Fisher', d1)])

# ---- Data 2
d2 = pd.read_csv(f'{MS}/Supplementary_Data_2_SMR_coloc.tsv', sep='\t')
c2 = [('context', CTX), ('region', 'Brain region.'), ('cell_type', 'Deconvolved cell type.'), ('disease', 'Disorder abbreviation. ' + DIS_FULL),
      ('GWAS_accession', 'GWAS Catalog accession of the summary statistics (Supplementary Table 2).'),
      ('gene_id', 'Ensembl gene ID (release 82).'), ('symbol', 'Gene symbol.'), ('biotype', 'Ensembl gene biotype.'),
      ('gene_chr', 'Chromosome of the gene.'), ('gene_start', 'Gene start (hg38).'),
      ('gene_in_MHC', 'TRUE when the gene start lies in chr6:25-35 Mb (hg38).'),
      ('top_SNP', 'Top cis-eQTL variant used by SMR.'), ('topSNP_chr', 'Chromosome of the top variant.'), ('topSNP_bp', 'Position of the top variant (hg38).'),
      ('p_GWAS', 'GWAS P of the top variant.'), ('p_eQTL', 'eQTL P of the top variant in the context.'),
      ('b_SMR', 'SMR effect (allele orientation corrected, see Methods).'), ('se_SMR', 'Standard error of b_SMR.'),
      ('p_SMR', 'SMR P.'), ('q_SMR', 'Benjamini-Hochberg q within the GWAS-context analysis; all rows have q_SMR < 0.05.'),
      ('p_HEIDI', 'HEIDI P; >= 0.01 indicates no evidence of heterogeneity.'), ('nsnp_HEIDI', 'Variants used by HEIDI.'),
      ('coloc_shared_SNPs', 'Variants shared by the eQTL and GWAS data in coloc.abf.'),
      ('PP_H3', 'Posterior probability of two distinct causal variants.'), ('PP_H4', 'Posterior probability of a shared causal variant; >= 0.8 taken as colocalization.'),
      ('all_three', 'TRUE when q_SMR < 0.05, p_HEIDI >= 0.01 and PP_H4 >= 0.8.'),
      ('low_proportion_context', 'TRUE for the 30 contexts whose median per-sample estimated proportion was below 0.001.')]
write(f'{out}/Supplementary_Data_2_v{ver}.xlsx', 'Supplementary Data 2. SMR, HEIDI and colocalization results',
      ['One row per context-disease-gene combination significant in SMR (q_SMR < 0.05), 2,578 rows covering 253 disease-gene pairs. ' + CTX],
      [('SMR_coloc', c2)], [('SMR_coloc', d2)])

# ---- Data 3
d3 = pd.read_csv(f'{SV}/Supplementary_Data_3_scGWAS_modules.tsv', sep='\t')
d3['disease'] = d3.disease.replace(DCODE)
c3 = [('context', CTX), ('disease', 'Disorder abbreviation. ' + DIS_FULL), ('gwas_accession', 'GWAS Catalog accession (Supplementary Table 2).'),
      ('seed_gene', 'Ensembl ID of the seed gene from which scGWAS grew the module.'), ('seed_symbol', 'Symbol of the seed gene.'),
      ('n_module_genes', 'Genes in the module.'), ('module_genes_ensembl', 'Module genes, Ensembl IDs separated by semicolons.'),
      ('module_genes_symbol', 'Module genes, symbols in the same order.'),
      ('reference_cell_type', 'Cell type of the PsychENCODE adult brain single-cell reference (DER-22) whose expression gave the expression score.'),
      ('module_score', 'Combined module score.'), ('gwas_score', 'GWAS component of the module score.'), ('expression_score', 'Expression component of the module score.'),
      ('module_score_z', 'Combined score normalized against size-matched random modules.'), ('gwas_score_z', 'Normalized GWAS score.'),
      ('expression_score_z', 'Normalized expression score.'),
      ('p_module', 'Empirical P of the combined score against size-matched random modules (0, no random module scored higher).'),
      ('p_gwas', 'Empirical P of the GWAS score.'), ('p_expression', 'Empirical P of the expression score.')]
write(f'{out}/Supplementary_Data_3_v{ver}.xlsx', 'Supplementary Data 3. Significant scGWAS modules',
      ['One row per module significant for the combined, GWAS and expression scores (each empirical P < 0.05): 118 modules in 35 context-GWAS combinations of eight disorders. '
       'Gene-level association scores came from MAGMA with SNP-gene annotations built from the eSNPs of each context. ' + CTX],
      [('scGWAS_modules', c3)], [('scGWAS_modules', d3)])

# ---- Data 4
d4a = pd.read_csv(f'{SV}/Supplementary_Data_4_external_SMR_pairs.tsv', sep='\t').rename(
    columns={'any_replicating_dataset_p_HEIDI_ge_0.01': 'any_dataset_p_HEIDI_ge_0.01'})
d4b = pd.read_csv(f'{SV}/Supplementary_Data_4_external_SMR_datasets.tsv', sep='\t')
for d in (d4a, d4b):
    d['disease'] = d.disease.replace(DCODE)
c4a = [('disease', 'Disorder abbreviation. ' + DIS_FULL), ('gene_ensembl', 'Ensembl gene ID.'), ('symbol', 'Gene symbol.'),
       ('passed_SMR_HEIDI_coloc', 'TRUE for the 32 pairs that passed SMR, HEIDI and colocalization in this study; FALSE for the 221 other SMR pairs.'),
       ('n_datasets_tested', 'External datasets (BrainMeta v2 and 13 GTEx v8 brain tissues) with a probe for the gene; 0, not testable.'),
       ('n_datasets_replicated', 'External datasets with P_SMR < 0.05/32 and the same sign of the SMR effect as in this study.'),
       ('best_external_p_SMR', 'Smallest external P_SMR.'), ('brainmeta_p_SMR', 'P_SMR in BrainMeta v2.'),
       ('any_dataset_p_HEIDI_ge_0.01', 'TRUE when at least one tested external dataset had P_HEIDI >= 0.01.'),
       ('replicated', 'TRUE when the pair replicated in at least one external dataset.')]
c4b = [('disease', 'Disorder abbreviation.'), ('gene_ensembl', 'Ensembl gene ID.'), ('symbol', 'Gene symbol.'),
       ('dataset', 'External eQTL dataset (BrainMeta v2, or a GTEx v8 brain tissue).'), ('topSNP', 'Top cis-eQTL variant in the external dataset.'),
       ('A1', 'Effect allele.'), ('A2', 'Other allele.'), ('Freq', 'Frequency of A1 in the LD reference.'),
       ('b_GWAS', 'GWAS effect of the top variant.'), ('se_GWAS', 'Standard error of b_GWAS.'), ('p_GWAS', 'GWAS P.'),
       ('b_eQTL', 'External eQTL effect.'), ('se_eQTL', 'Standard error of b_eQTL.'), ('p_eQTL', 'External eQTL P.'),
       ('b_SMR', 'External SMR effect.'), ('se_SMR', 'Standard error of b_SMR.'), ('p_SMR', 'External SMR P.'),
       ('p_HEIDI', 'External HEIDI P.'), ('nsnp_HEIDI', 'Variants used by HEIDI.'),
       ('b_SMR_this_study', 'SMR effect in this study, from the context with the smallest P_SMR for the pair.'),
       ('p_SMR_this_study', 'SMR P in that context.'), ('same_direction', 'TRUE when b_SMR and b_SMR_this_study have the same sign.'),
       ('replicated', 'TRUE when p_SMR < 0.05/32 and same_direction.')]
write(f'{out}/Supplementary_Data_4_v{ver}.xlsx', 'Supplementary Data 4. Replication of SMR pairs in external brain eQTL data',
      ['SMR was repeated for the 253 disease-gene pairs significant in this study with BrainMeta v2 cortical eQTLs and 13 GTEx v8 brain tissues, with the GWAS files, LD reference and settings of the main analysis.',
       'Sheet "Pairs": one row per disease-gene pair. Sheet "Datasets": one row per pair and external dataset with an SMR result.'],
      [('Pairs', c4a), ('Datasets', c4b)], [('Pairs', d4a), ('Datasets', d4b)])
print('XLSX_DONE')
