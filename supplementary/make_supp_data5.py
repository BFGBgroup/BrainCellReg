# -*- coding: utf-8 -*-
"""Supplementary Data 5: positional and functional features of eSNPs.
Tables behind Fig. 5 and Supplementary Figs. 6-7, as computed on the cluster (esnp_eg/results). Only the eSNP set used in
the paper (eset == "all") is kept; columns are renamed, no number is changed. Set SRC to the folder with the result tables.
python make_supp_data5.py <out.xlsx>"""
import sys, io, os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
out = sys.argv[1]
SRC = r'<path to esnp_eg/results>'
LEVEL = {'WB': 'Whole brain', 'cell': 'Cell type', 'region': 'Region', 'ctx': 'Context'}


def rd(name):
    return pd.read_csv(os.path.join(SRC, name), sep='\t')


def split_group(df, col='group'):
    g = df[col].astype(str)
    lvl = g.map(lambda x: LEVEL['WB'] if x == 'WB' else LEVEL[x.split(':', 1)[0]])
    nm = g.map(lambda x: 'Whole brain' if x == 'WB' else x.split(':', 1)[1])
    df = df.drop(columns=[col])
    df.insert(0, 'group', nm)
    df.insert(0, 'level', lvl)
    return df


def all_set(df):
    return df[df['eset'] == 'all'].drop(columns=['eset']).reset_index(drop=True)


def write(path, title, notes, cols, sheets):   # same layout as make_supp_xlsx.py (Description sheet first)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Description'
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=12)
    r = 3
    for n in notes:
        ws.cell(r, 1, n).alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
    r += 1
    for sheet, cdesc in cols:
        ws.cell(r, 1, f'Columns of sheet "{sheet}"').font = Font(bold=True)
        r += 1
        for c, d in cdesc:
            ws.cell(r, 1, c).font = Font(name='Consolas')
            ws.cell(r, 2, d).alignment = Alignment(wrap_text=True, vertical='top')
            r += 1
        r += 1
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 110
    for name, df in sheets:
        s = wb.create_sheet(name)
        s.append(list(df.columns))
        for c in s[1]:
            c.font = Font(bold=True)
        for row in df.itertuples(index=False):
            s.append([None if (isinstance(v, float) and pd.isna(v)) else (v.item() if hasattr(v, 'item') else v) for v in row])
        s.freeze_panes = 'A2'
        for i, c in enumerate(df.columns, 1):
            s.column_dimensions[get_column_letter(i)].width = min(45, max(10, len(str(c)) + 2))
        documented = [c for c, _ in dict(cols)[name]]
        assert documented == list(df.columns), (name, set(df.columns) ^ set(documented))
    wb.save(path)


LV = [('level', 'Whole brain (union of eSNPs across contexts), Cell type, Region or Context'),
      ('group', 'Cell type, region, or context written REGION.CellType (for example, FCTX.ExN)')]
FISH = [('n_esnp', 'Independent eSNPs in the group (after LD clumping)'), ('n_non', 'Non-eSNPs in the group (after LD clumping)'),
        ('esnp_in', 'eSNPs in the category'), ('non_in', 'Non-eSNPs in the category')]
ORC = [('OR', 'Odds ratio, eSNPs against non-eSNPs (two-sided Fisher\'s exact test)'), ('OR_lo', 'Lower 95% confidence limit of the OR'),
       ('OR_hi', 'Upper 95% confidence limit of the OR'), ('p', 'Two-sided P value')]
PCT = [('pct_esnp', 'Percentage of eSNPs in the category'), ('pct_non', 'Percentage of non-eSNPs in the category')]
Q = [('q', 'Benjamini-Hochberg q value within the analysis')]

# region classes
rc = split_group(all_set(rd('step7_region_class.tsv')))
rc = rc.rename(columns={'feature': 'region_class'})
rc_cols = LV + [('region_class', 'ANNOVAR region class')] + FISH + PCT + ORC + Q
rc = rc[[c for c, _ in rc_cols]]
# TSS
tb = rd('step6_tss_bins.tsv').rename(columns={'scope': 'group_raw', 'bin': 'bin_10kb'})
tb = split_group(tb, 'group_raw')
tb_cols = LV[:1] + [('group', 'Whole brain, cell type or region')] + [
    ('bin_10kb', '10-kb bin of distance from the TSS of the gene (negative, upstream; strand-aware), -100 to 100 = -1 to 1 Mb'),
    ('n_tested', 'Tested variant-gene pairs in the bin'), ('n_esnp_pairs', 'eSNP-eGene pairs in the bin'),
    ('n_esnp_pairs_lnc', 'eSNP-eGene pairs of lncRNA eGenes'), ('n_esnp_pairs_pc', 'eSNP-eGene pairs of protein-coding eGenes'),
    ('pct_esnp', 'eSNP-eGene pairs as a percentage of tested pairs in the bin')]
tb = tb[[c for c, _ in tb_cols]]
tc = rd('tss_by_ctx.tsv').rename(columns={'ctx': 'context'})
tc_cols = [('context', 'Context written REGION.CellType'), ('n_pairs', 'eSNP-eGene pairs'),
           ('pct_10kb', 'Percentage of eSNP-eGene pairs within 10 kb of the TSS'), ('pct_100kb', 'Percentage within 100 kb of the TSS'),
           ('median_abs_dist', 'Median absolute distance to the TSS (bp)')]
# MAF
mb = split_group(all_set(rd('maf_bins.tsv')))
mb_cols = LV + [('maf_bin', 'Minor allele frequency bin'), ('n_cis', 'Tested variants in the bin'),
                ('n_esnp', 'Independent eSNPs in the group'), ('n_non', 'Non-eSNPs in the group'),
                ('prop_esnp_in_cis', 'Proportion of eSNPs in the bin'), ('prop_non_in_cis', 'Proportion of non-eSNPs in the bin')] + ORC
mb = mb[[c for c, _ in mb_cols]]
mk = split_group(rd('maf_ks_by_group.tsv'))
mk_cols = LV + [('n_esnp', 'Independent eSNPs'), ('n_non', 'Non-eSNPs'), ('median_esnp', 'Median MAF of eSNPs'),
                ('median_non', 'Median MAF of non-eSNPs'), ('D', 'Kolmogorov-Smirnov D'), ('p', 'Two-sided Kolmogorov-Smirnov P value')]
# chromatin, DNase and brain TF features by group (incl. every context)
cf = split_group(rd('ctx_features_extra.tsv'))
cf_cols = LV + [('feature', 'Prefrontal cortex histone mark or DNase peaks (ENCODE), ReMap 2022 binding sites, or continuous score'),
                ('test', 'fisher, overlap of eSNPs against non-eSNPs; ks, Kolmogorov-Smirnov test of a continuous score'),
                ('n1', 'eSNPs tested'), ('n2', 'Non-eSNPs tested'), ('a', 'eSNPs in the feature (Fisher tests)'),
                ('c', 'Non-eSNPs in the feature (Fisher tests)'), ('OR', 'Odds ratio, eSNPs against non-eSNPs (Fisher tests)'),
                ('OR_lo', 'Lower 95% confidence limit of the OR'), ('OR_hi', 'Upper 95% confidence limit of the OR'),
                ('p', 'Two-sided P value'), ('D', 'Kolmogorov-Smirnov D (ks tests)'),
                ('direction', 'Direction of the difference, eSNPs relative to non-eSNPs')]
cf = cf[[c for c, _ in cf_cols]]
# cCRE
cc = split_group(all_set(rd('step10_ccre.tsv'))).rename(columns={'feature': 'cCRE_class'})
cc_cols = LV + [('cCRE_class', 'ENCODE cCRE class (PLS, promoter-like; pELS / dELS, proximal / distal enhancer-like)')] + FISH + PCT + ORC + Q
cc = cc[[c for c, _ in cc_cols]]
# ReMap
rt = split_group(all_set(rd('step10_remap_tf.tsv'))).rename(columns={'feature': 'TF'})
rt_cols = LV[:1] + [('group', 'Whole brain, cell type or region')] + [('TF', 'ReMap 2022 transcription factor (peaks of all biosamples merged); any_TF_peak, any peak')] + FISH + ORC + PCT + Q
rt = rt[[c for c, _ in rt_cols]]
rb = split_group(all_set(rd('step10_remap_brain_biotypes.tsv'))).rename(columns={'feature': 'TF_biosample'})
rb_cols = LV[:1] + [('group', 'Whole brain, cell type or region')] + [('TF_biosample', 'ReMap 2022 transcription factor and brain or neural biosample')] + FISH + PCT + ORC + Q
rb = rb[[c for c, _ in rb_cols]]
# Enformer, FAVOR
SC = [('n_esnp', 'Independent eSNPs'), ('n_non', 'Non-eSNPs'), ('median_esnp', 'Median score of eSNPs'), ('median_non', 'Median score of non-eSNPs'),
      ('mean_esnp', 'Mean score of eSNPs'), ('mean_non', 'Mean score of non-eSNPs'), ('auc', 'Area under the ROC curve separating eSNPs from non-eSNPs'),
      ('p', 'Two-sided P value'), ('q', 'Benjamini-Hochberg q value within the analysis')]
ef = split_group(all_set(rd('step8_enformer.tsv'))).rename(columns={'feature': 'score'})
ef_cols = LV[:1] + [('group', 'Whole brain, cell type or region')] + [('score', 'Enformer principal component (PC0-PC19) or L2 norm of the 20 PCs')] + SC
ef = ef[[c for c, _ in ef_cols]]
fv = split_group(all_set(rd('step9_favor_apc.tsv'))).rename(columns={'feature': 'score'})
fv_cols = LV[:1] + [('group', 'Whole brain, cell type or region')] + [('score', 'FAVOR annotation principal component (apc_*), CADD PHRED or LINSIGHT')] + SC
fv = fv[[c for c, _ in fv_cols]]
ks1 = split_group(rd('ks_by_group.tsv').drop(columns=['level'])).rename(columns={'feature': 'score'})   # source 'level' repeats group
ks2 = split_group(rd('ks_pcs_by_group.tsv')).rename(columns={'feature': 'score'})
KS = [('n_esnp', 'Independent eSNPs'), ('n_non', 'Non-eSNPs'), ('median_esnp', 'Median score of eSNPs'), ('median_non', 'Median score of non-eSNPs'),
      ('D', 'Kolmogorov-Smirnov D'), ('direction', 'Direction, eSNPs relative to non-eSNPs'), ('p', 'Two-sided Kolmogorov-Smirnov P value')]
ks1_cols = LV + [('score', 'FAVOR aPC or Enformer L2 norm')] + KS + [('q', 'Benjamini-Hochberg q value')]
ks1 = ks1[[c for c, _ in ks1_cols]]
ks2_cols = LV[:1] + [('group', 'Whole brain, cell type or region')] + [('score', 'Enformer principal component')] + KS
ks2 = ks2[[c for c, _ in ks2_cols]]

notes = [
    'Positional and functional features of eSNPs (Fig. 5 and Supplementary Figs. 6 and 7). In each group, independent eSNPs '
    '(after LD clumping) were compared with non-eSNPs, the tested variants whose nominal P was at least 0.05 for every gene in the '
    'same context, also after LD clumping. Whole brain is the union of eSNPs across the 110 contexts.',
    'Categorical features (region class, MAF bin, histone marks, DNase peaks, cCREs, transcription factor binding sites) were tested '
    'with two-sided Fisher\'s exact tests; continuous scores (MAF, Enformer, FAVOR) with two-sided Kolmogorov-Smirnov tests and the AUC.',
    'Histone marks and DNase peaks are ENCODE prefrontal cortex peak sets; cCREs are the ENCODE registry; transcription factor binding '
    'sites are from ReMap 2022.',
]
cols = [('Region class', rc_cols), ('TSS bins', tb_cols), ('TSS by context', tc_cols), ('MAF bins', mb_cols), ('MAF KS', mk_cols),
        ('Chromatin by context', cf_cols), ('cCRE', cc_cols), ('ReMap TF', rt_cols), ('ReMap brain', rb_cols),
        ('Enformer', ef_cols), ('FAVOR', fv_cols), ('KS scores', ks1_cols), ('KS Enformer PCs', ks2_cols)]
sheets = [('Region class', rc), ('TSS bins', tb), ('TSS by context', tc), ('MAF bins', mb), ('MAF KS', mk),
          ('Chromatin by context', cf), ('cCRE', cc), ('ReMap TF', rt), ('ReMap brain', rb),
          ('Enformer', ef), ('FAVOR', fv), ('KS scores', ks1), ('KS Enformer PCs', ks2)]
write(out, 'Supplementary Data 5. Positional and functional features of eSNPs', notes, cols, sheets)
for n, d in sheets:
    print(f'{n}: {len(d)} rows')
print('wrote', out)
